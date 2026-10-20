import io
import zipfile

import pandas as pd
from openpyxl import load_workbook

from enumerator_reports import enumerator_error_zip


def test_error_log_zip_contains_one_formatted_xlsx_per_enumerator():
    source = pd.DataFrame({
        "Collector Name": ["Ana", "Beto"],
        "Quadrat Photo_URL": [
            "https://example.test/ana-photo",
            "https://example.test/beto-photo",
        ],
    })
    issues = pd.DataFrame([
        {
            "row_index": 0, "uuid": "u1", "category": "gps", "field": "_Local GPS",
            "severity": "error", "message": "Missing GPS coordinates",
        },
        {
            "row_index": 1, "uuid": "u2", "category": "schema", "field": "schema",
            "severity": "error", "message": "Unsupported Kobo form schema",
        },
    ])

    payload = enumerator_error_zip(source, issues)

    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        names = archive.namelist()
        assert names == ["01_Ana_errors.xlsx", "02_Beto_errors.xlsx"]
        ana_bytes = archive.read(names[0])

    workbook = load_workbook(io.BytesIO(ana_bytes), data_only=True)
    sheet = workbook["Findings"]
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == sheet.dimensions
    headers = [cell.value for cell in sheet[1]]
    assert "Finding / suggested action" in headers
    message_column = headers.index("Finding / suggested action") + 1
    assert sheet.cell(2, message_column).value == (
        "Koordenada GPS falta (Missing GPS coordinates)"
    )
    photo_column = headers.index("Kobo photo URL") + 1
    assert sheet.cell(2, photo_column).hyperlink.target == "https://example.test/ana-photo"


def test_error_log_zip_keeps_technical_messages_in_english():
    source = pd.DataFrame({"Collector Name": ["Ana"]})
    issues = pd.DataFrame([{
        "row_index": 0, "uuid": "u1", "category": "schema", "field": "schema",
        "severity": "error", "message": "Unsupported Kobo form schema",
    }])

    payload = enumerator_error_zip(source, issues)
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        workbook = load_workbook(io.BytesIO(archive.read(archive.namelist()[0]), data_only=True)
    sheet = workbook["Findings"]
    headers = [cell.value for cell in sheet[1]]
    assert sheet.cell(2, headers.index("Finding / suggested action") + 1).value == (
        "Unsupported Kobo form schema"
    )
