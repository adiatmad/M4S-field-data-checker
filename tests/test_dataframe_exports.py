import io
from datetime import datetime, timezone, timedelta

import pandas as pd
from openpyxl import load_workbook

from dataframe_exports import dataframe_to_xlsx


def test_xlsx_export_preserves_timezone_aware_timestamps_as_iso_text():
    frame = pd.DataFrame({
        "submission_time": pd.to_datetime([
            "2026-10-20T06:30:00+00:00",
            "2026-10-20T13:30:00+07:00",
        ], utc=True),
    })

    workbook = load_workbook(io.BytesIO(dataframe_to_xlsx(frame)), data_only=True)
    sheet = workbook["Data"]

    assert sheet["A2"].value == "2026-10-20T06:30:00+00:00"
    assert sheet["A3"].value == "2026-10-20T06:30:00+00:00"


def test_xlsx_export_handles_timezone_values_inside_mixed_object_column():
    frame = pd.DataFrame({
        "identifier": [670, "_670__76394171"],
        "timestamp": [
            datetime(2026, 10, 20, 13, 0, tzinfo=timezone(timedelta(hours=7))),
            "not a timestamp",
        ],
    }, dtype=object)

    workbook = load_workbook(io.BytesIO(dataframe_to_xlsx(frame)), data_only=True)
    sheet = workbook["Data"]

    assert sheet["A2"].value == 670
    assert sheet["A3"].value == "_670__76394171"
    assert sheet["B2"].value == "2026-10-20T13:00:00+07:00"
    assert sheet["B3"].value == "not a timestamp"


def test_xlsx_export_keeps_filters_and_frozen_header():
    frame = pd.DataFrame({"identifier": ["a"], "value": [1]})

    workbook = load_workbook(io.BytesIO(dataframe_to_xlsx(frame, sheet_name="QA issues")))
    sheet = workbook["QA issues"]

    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:B2"
