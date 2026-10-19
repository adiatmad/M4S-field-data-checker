import io
import zipfile

import pandas as pd
from openpyxl import load_workbook

from enumerator_reports import (
    build_enumerator_raw_data_reviews,
    enumerator_raw_data_review_workbook,
    enumerator_raw_data_review_zip,
)


def _fixture():
    raw = pd.DataFrame({
        "Naran Koletor": ["Enumerator A", "Enumerator B"],
        "Espesie du'ut tasi (seagrass) ne'ebe prezensa iha Quadrante/Halophila ovalis": [1, 0],
        "Persentajen (Ho) Halophila ovalis  (%)": [None, 0],
        "Quadrat Photo_URL": ["https://example.test/photo-a", "https://example.test/photo-b"],
        "Unrelated raw field": ["keep-a", "keep-b"],
    })
    issues = pd.DataFrame([
        {
            "row_index": 0, "uuid": "uuid-a", "category": "species",
            "field": "Percent Halophila ovalis (%)", "severity": "error",
            "message": "Halophila ovalis marked present but no % cover recorded",
        },
    ])
    return raw, issues


def test_review_rows_preserve_all_raw_columns_and_append_qa_message():
    raw, issues = _fixture()
    reviews = build_enumerator_raw_data_reviews(raw, issues)

    assert list(reviews) == ["Enumerator A"]
    frame = reviews["Enumerator A"]
    assert list(raw.columns).issubset(frame.columns)
    assert frame.iloc[0]["Persentajen (Ho) Halophila ovalis  (%)"] is None or pd.isna(
        frame.iloc[0]["Persentajen (Ho) Halophila ovalis  (%)"]
    )
    assert frame.iloc[0]["Unrelated raw field"] == "keep-a"
    assert frame.iloc[0]["QA finding count"] == 1
    assert "no % cover recorded" in frame.iloc[0]["QA message"]
    assert frame.iloc[0]["QA fields to review"] == "Percent Halophila ovalis (%)"


def test_review_workbook_contains_original_ho_cover_column_and_per_enumerator_sheet():
    raw, issues = _fixture()
    payload = enumerator_raw_data_review_workbook(raw, issues)
    workbook = load_workbook(io.BytesIO(payload), read_only=True)

    assert workbook.sheetnames == ["Summary", "Enumerator A"]
    headers = [cell.value for cell in next(workbook["Enumerator A"].iter_rows(min_row=1, max_row=1))]
    assert "Persentajen (Ho) Halophila ovalis  (%)" in headers
    assert "QA message" in headers
    assert "Unrelated raw field" in headers


def test_review_zip_keeps_original_columns_and_excel_friendly_csv():
    raw, issues = _fixture()
    payload = enumerator_raw_data_review_zip(raw, issues)
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        names = archive.namelist()
        assert len(names) == 1
        content = archive.read(names[0])
        assert content.startswith(b"\xef\xbb\xbf")
        assert b"Persentajen (Ho) Halophila ovalis  (%)" in content
        assert b"QA message" in content
        assert b"Unrelated raw field" in content
