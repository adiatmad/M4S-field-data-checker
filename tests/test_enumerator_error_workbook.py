import io

import pandas as pd
from openpyxl import load_workbook

from enumerator_reports import enumerator_error_zip, enumerator_error_workbook


def test_enumerator_error_workbook_is_formatted_and_filterable():
    source = pd.DataFrame(
        {
            "Collector Name": ["Enumerator A", "Enumerator B"],
            "Quadrat Photo_URL": [
                "https://example.org/photo-a.jpg",
                "https://example.org/photo-b.jpg",
            ],
        },
        index=[0, 1],
    )
    issues = pd.DataFrame(
        [
            {
                "row_index": 0,
                "uuid": "uuid-a",
                "category": "photo",
                "field": "Quadrat Photo_URL",
                "severity": "error",
                "message": "No quadrat photo attached",
            },
            {
                "row_index": 1,
                "uuid": "uuid-b",
                "category": "gps",
                "field": "_Local GPS",
                "severity": "warning",
                "message": "Verify coordinates",
            },
        ]
    )

    workbook_bytes = enumerator_error_workbook(source, issues)
    workbook = load_workbook(io.BytesIO(workbook_bytes))

    assert workbook.sheetnames == ["Summary", "Enumerator A", "Enumerator B"]
    summary = workbook["Summary"]
    assert [cell.value for cell in summary[1]] == [
        "Enumerator", "Findings", "Errors", "Warnings"
    ]
    assert [summary.cell(row=2, column=col).value for col in range(1, 5)] == [
        "Enumerator A", 1, 1, 0
    ]

    sheet = workbook["Enumerator A"]
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == sheet.dimensions
    assert [cell.value for cell in sheet[1]] == [
        "Source row",
        "Submission UUID",
        "Severity",
        "Issue category",
        "Field to check",
        "Finding / suggested action",
        "Kobo photo URL",
    ]
    assert sheet["C2"].value == "error"
    assert sheet["C2"].fill.fgColor.rgb.endswith("FCE4D6")
    assert sheet["G2"].hyperlink.target == "https://example.org/photo-a.jpg"


def test_enumerator_error_workbook_handles_no_findings():
    workbook_bytes = enumerator_error_workbook(
        pd.DataFrame({"Collector Name": ["Enumerator A"]}),
        pd.DataFrame(columns=["row_index", "uuid", "category", "field", "severity", "message"]),
    )
    workbook = load_workbook(io.BytesIO(workbook_bytes))
    assert workbook.sheetnames == ["Summary"]
    assert workbook["Summary"]["A2"].value == "No findings for the selected report options"



def test_enumerator_error_zip_contains_directly_openable_xlsx():
    import zipfile

    source = pd.DataFrame({"Collector Name": ["Enumerator A"]}, index=[0])
    message = "Check this value, then confirm the field's meaning."
    issues = pd.DataFrame([{
        "row_index": 0,
        "uuid": "uuid-a",
        "category": "text",
        "field": "Local notes",
        "severity": "warning",
        "message": message,
    }])

    archive_bytes = enumerator_error_zip(source, issues)
    with zipfile.ZipFile(io.BytesIO(archive_bytes)) as archive:
        names = archive.namelist()
        assert names == ["01_Enumerator_A_errors.xlsx"]
        workbook = load_workbook(io.BytesIO(archive.read(names[0])), data_only=True)

    sheet = workbook["Findings"]
    headers = [cell.value for cell in sheet[1]]
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == sheet.dimensions
    assert sheet.cell(2, headers.index("Finding / suggested action") + 1).value == message
    assert sheet.cell(2, headers.index("Source row") + 1).value == 0
