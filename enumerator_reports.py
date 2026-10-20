"""Enumerator QA report exports, kept separate from Streamlit page execution."""

import io
import re
import pandas as pd

from qa_message_localization import localize_qa_issue_messages


def _issue_export_frame(df, issues):
    """Enrich findings with enumerator and direct Kobo image URL for field review."""
    enriched = issues.copy()
    if "enumerator" not in enriched.columns:
        enriched["enumerator"] = ""
    if "photo_url" not in enriched.columns:
        enriched["photo_url"] = ""
    if df is None or df.empty:
        return enriched
    enumerator_col = "Collector Name" if "Collector Name" in df.columns else None
    photo_col = "Quadrat Photo_URL" if "Quadrat Photo_URL" in df.columns else None
    for issue_index, issue in enriched.iterrows():
        row_index = issue.get("row_index")
        try:
            source = df.loc[row_index]
            if isinstance(source, pd.DataFrame):
                source = source.iloc[0]
        except (KeyError, TypeError, IndexError):
            continue
        if enumerator_col:
            enriched.at[issue_index, "enumerator"] = source.get(enumerator_col, "")
        if photo_col:
            value = source.get(photo_col, "")
            enriched.at[issue_index, "photo_url"] = "" if pd.isna(value) else str(value).strip()
    return enriched


def filter_enumerator_report_issues(issues, *, include_gps_accuracy=True, include_site_distance=True,
                                      include_text_formatting=True):
    """Filter optional GPS/text-cleaning findings from enumerator reports only."""
    visible = issues.copy()
    if not include_text_formatting and "category" in visible.columns:
        visible = visible.loc[visible["category"] != "text_formatting"]
    if not include_gps_accuracy and "field" in visible.columns:
        precision = visible["field"].astype("string").str.contains("precision", case=False, na=False)
        visible = visible.loc[~((visible["category"] == "gps") & precision)]
    if not include_site_distance and "message" in visible.columns:
        distance = visible["message"].astype("string").str.contains(
            "centroid-distance threshold", case=False, na=False
        )
        visible = visible.loc[~((visible["category"] == "gps") & distance)]
    return visible.reset_index(drop=True)


def build_enumerator_error_logs(df, issues, *, include_gps_accuracy=True, include_site_distance=True,
                               include_text_formatting=True):
    """Return one enriched CSV dataframe per enumerator, sorted by enumerator name."""
    filtered = filter_enumerator_report_issues(
        issues, include_gps_accuracy=include_gps_accuracy,
        include_site_distance=include_site_distance,
        include_text_formatting=include_text_formatting,
    )
    # Translate recognized survey errors only in this presentation/export copy.
    filtered = localize_qa_issue_messages(filtered)
    enriched = _issue_export_frame(df, filtered)
    if enriched.empty:
        return {}
    enriched["enumerator"] = enriched["enumerator"].fillna("").astype(str).str.strip()
    enriched.loc[enriched["enumerator"].eq(""), "enumerator"] = "Unknown enumerator"
    logs = {}
    for name in sorted(enriched["enumerator"].unique(), key=str.casefold):
        logs[name] = enriched.loc[enriched["enumerator"] == name].sort_values(
            ["row_index", "severity"], kind="stable", na_position="last"
        ).reset_index(drop=True)
    return logs


def enumerator_error_zip(df, issues, *, include_gps_accuracy=True, include_site_distance=True,
                        include_text_formatting=True):
    """Build a ZIP of individual enumerator CSV error logs."""
    import io
    import zipfile

    logs = build_enumerator_error_logs(
        df, issues, include_gps_accuracy=include_gps_accuracy,
        include_site_distance=include_site_distance,
        include_text_formatting=include_text_formatting,
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for index, (name, frame) in enumerate(logs.items(), start=1):
            safe = re.sub(r"[^A-Za-z0-9._-]+", "_", name).strip("._") or "enumerator"
            csv_text = frame.to_csv(index=False, sep=";")
            # Excel on Windows often uses the regional list separator for double-click CSV opening.
            # A UTF-8 BOM preserves non-ASCII field text in Excel.
            archive.writestr(f"{index:02d}_{safe}_errors.csv", "\ufeff" + csv_text)
    return buffer.getvalue()


def enumerator_error_workbook(df, issues, *, include_gps_accuracy=True, include_site_distance=True,
                             include_text_formatting=True):
    """Build one Excel workbook with a summary and one findings sheet per enumerator."""
    logs = build_enumerator_error_logs(
        df, issues, include_gps_accuracy=include_gps_accuracy,
        include_site_distance=include_site_distance,
        include_text_formatting=include_text_formatting,
    )
    summary_rows = []
    for name, frame in logs.items():
        severity = frame.get("severity", pd.Series(dtype="string")).astype("string").str.casefold()
        summary_rows.append({
            "Enumerator": name,
            "Findings": len(frame),
            "Errors": int((severity == "error").sum()),
            "Warnings": int((severity == "warning").sum()),
        })
    summary = pd.DataFrame(summary_rows, columns=["Enumerator", "Findings", "Errors", "Warnings"])
    if summary.empty:
        summary = pd.DataFrame([{
            "Enumerator": "No findings for the selected report options",
            "Findings": 0, "Errors": 0, "Warnings": 0,
        }])

    # Present findings in a consistent, review-friendly order without modifying source data.
    display_columns = {
        "row_index": "Source row",
        "uuid": "Submission UUID",
        "category": "Issue category",
        "field": "Field to check",
        "severity": "Severity",
        "message": "Finding / suggested action",
        "photo_url": "Kobo photo URL",
    }
    preferred = ["row_index", "uuid", "severity", "category", "field", "message", "photo_url"]

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        used_names = {"summary"}
        for name, frame in logs.items():
            base = str(name).strip("' ") or "Enumerator"
            for forbidden in "[]:*?/\\\\":
                base = base.replace(forbidden, "_")
            candidate = base[:31]
            suffix = 2
            while candidate.casefold() in used_names:
                marker = f"_{suffix}"
                candidate = f"{base[:31-len(marker)]}{marker}"
                suffix += 1
            used_names.add(candidate.casefold())
            ordered = [column for column in preferred if column in frame.columns]
            # Retain any future fields instead of silently dropping them.
            ordered += [column for column in frame.columns if column not in ordered and column != "enumerator"]
            sheet = frame.loc[:, ordered].rename(columns=display_columns)
            sheet.to_excel(writer, sheet_name=candidate, index=False)

        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        header_fill = PatternFill("solid", fgColor="1F4E78")
        header_font = Font(color="FFFFFF", bold=True)
        for worksheet in writer.book.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            worksheet.sheet_view.showGridLines = False
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            for column_cells in worksheet.columns:
                column_letter = get_column_letter(column_cells[0].column)
                header = str(column_cells[0].value or "")
                values = [len(str(cell.value or "")) for cell in column_cells[:200]]
                width = min(max(max(values, default=10) + 2, 12), 52)
                if header in ("Finding / suggested action", "Kobo photo URL"):
                    width = 48
                elif header in ("Submission UUID", "Enumerator"):
                    width = 28
                worksheet.column_dimensions[column_letter].width = width
            for row in worksheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=cell.column == 6)
                if worksheet.title != "Summary":
                    headers = {cell.value: cell.column for cell in worksheet[1]}
                    severity_column = headers.get("Severity")
                    if severity_column:
                        severity_cell = row[severity_column - 1]
                        severity = str(severity_cell.value or "").casefold()
                        if severity == "error":
                            severity_cell.fill = PatternFill("solid", fgColor="FCE4D6")
                            severity_cell.font = Font(color="9C0006", bold=True)
                        elif severity == "warning":
                            severity_cell.fill = PatternFill("solid", fgColor="FFF2CC")
                            severity_cell.font = Font(color="9C6500", bold=True)
                    photo_column = headers.get("Kobo photo URL")
                    if photo_column:
                        photo_cell = row[photo_column - 1]
                        url = str(photo_cell.value or "").strip()
                        if url.startswith(("https://", "http://")):
                            photo_cell.hyperlink = url
                            photo_cell.style = "Hyperlink"
    return buffer.getvalue()
